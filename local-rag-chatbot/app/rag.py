"""RAG 파이프라인: 텍스트 추출 → 청크 분할 → 하이브리드 검색.

검색은 두 가지를 합친다 (Reciprocal Rank Fusion).
  1) 키워드: 한국어 글자 2-gram BM25 — 형태소 분석기 없이도 조사·어미 변화에 강하다
  2) 의미: Ollama 임베딩(bge-m3) 코사인 유사도 — 표현이 달라도 뜻으로 찾는다
임베딩 서버가 없으면 키워드 검색만으로 동작한다.
"""
import io
import logging
import math
import os
import re
import threading
from collections import Counter

from . import ollama, store

log = logging.getLogger(__name__)

EMBED_MODEL = os.getenv("OLLAMA_EMBED_MODEL", "bge-m3").strip()


# ---------- 1. 텍스트 추출 ----------

def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if name.endswith((".txt", ".md", ".csv", ".tsv")):
        for enc in ("utf-8-sig", "cp949", "euc-kr"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="ignore")
    raise ValueError("지원하지 않는 형식입니다. (pdf, docx, txt, md, csv)")


# ---------- 2. 청크 분할 ----------

def chunk_text(text: str, size: int = 600, overlap: int = 120) -> list[str]:
    """문단 단위로 묶어 size 글자 안팎의 조각을 만들고, 앞 조각 끝 overlap 글자를 이어 붙인다."""
    text = re.sub(r"[ \t]+", " ", text)
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if len(cur) + len(p) + 1 <= size:
            cur = f"{cur}\n{p}".strip()
            continue
        if cur:
            chunks.append(cur)
            cur = cur[-overlap:]
        while len(cur) + len(p) > size:
            take = size - len(cur)
            chunks.append((cur + p[:take]).strip())
            cur, p = p[max(take - overlap, 0):take], p[take:]
        cur = f"{cur}{p}".strip()
    if cur:
        chunks.append(cur)
    return chunks


# ---------- 3-1. 키워드 검색 (BM25) ----------

# 질문·문장 어미에서 나오는 2-gram. 내용과 무관하게 FAQ류 문서와 겹쳐 순위를 흐리므로 뺀다.
# (eval.py 결과로 찾은 문제: "~하나요?" 질문이 FAQ 문서로 쏠림)
STOP_BIGRAMS = {
    "하나", "나요", "있나", "까요", "니까", "합니", "니다", "습니", "됩니", "인가", "가요",
    "어떻", "떻게", "하고", "싶어", "어요", "에요", "해요", "세요", "하는", "하면", "으면",
}


def tokenize(text: str) -> list[str]:
    toks = []
    for w in re.findall(r"[0-9a-z]+|[가-힣]+", text.lower()):
        if "가" <= w[0] <= "힣":
            if len(w) == 1:
                toks.append(w)
            else:
                toks.extend(b for b in (w[i:i + 2] for i in range(len(w) - 1)) if b not in STOP_BIGRAMS)
            if len(w) > 2:
                toks.append(w)  # 어절 전체가 같으면 가산점
        else:
            toks.append(w)
    return toks


class BM25:
    def __init__(self, docs: list[list[str]], k1=1.5, b=0.75):
        self.k1, self.b = k1, b
        self.tfs = [Counter(d) for d in docs]
        self.lens = [len(d) for d in docs]
        self.avg = (sum(self.lens) / len(docs)) if docs else 0
        df = Counter(t for d in docs for t in set(d))
        n = len(docs)
        self.idf = {t: math.log(1 + (n - f + 0.5) / (f + 0.5)) for t, f in df.items()}

    def scores(self, query: list[str]) -> list[float]:
        out = []
        for tf, ln in zip(self.tfs, self.lens):
            s = 0.0
            for t in query:
                if t in tf:
                    f = tf[t]
                    s += self.idf[t] * f * (self.k1 + 1) / (f + self.k1 * (1 - self.b + self.b * ln / self.avg))
            out.append(s)
        return out


# ---------- 3-2. 의미 검색 (임베딩) ----------

def embed(texts: list[str]):
    import numpy as np
    vecs = np.asarray(ollama.embed(EMBED_MODEL, texts), dtype="float32")
    norms = np.linalg.norm(vecs, axis=1, keepdims=True)
    return vecs / np.where(norms == 0, 1, norms)


# ---------- 4. 인덱스 ----------

class Index:
    def __init__(self):
        self.lock = threading.Lock()
        self.chunks: list[dict] = []
        self.bm25 = BM25([])
        self.matrix = None

    def rebuild(self):
        if store.get_meta("embed_model") != EMBED_MODEL:
            store.clear_embeddings()  # 임베딩 모델이 바뀌면 다시 계산
            store.set_meta("embed_model", EMBED_MODEL)
        chunks = store.all_chunks()
        matrix = None
        if EMBED_MODEL and chunks:
            try:
                matrix = self._dense_matrix(chunks)
            except Exception as e:
                log.warning("임베딩 검색을 쓸 수 없어 키워드 검색만 사용합니다: %s", e)
        bm25 = BM25([tokenize(f'{c["title"]} {c["text"]}') for c in chunks])
        with self.lock:
            self.chunks, self.bm25, self.matrix = chunks, bm25, matrix

    @staticmethod
    def _dense_matrix(chunks):
        import numpy as np
        missing = [c for c in chunks if c["embedding"] is None]
        if missing:
            log.info("임베딩 계산 중: %d개 조각", len(missing))
            for c, v in zip(missing, embed([c["text"] for c in missing])):
                c["embedding"] = v.tobytes()
                store.set_chunk_embedding(c["id"], c["embedding"])
        return np.vstack([np.frombuffer(c["embedding"], dtype="float32") for c in chunks])

    def search(self, query: str, k: int = 5) -> list[dict]:
        with self.lock:
            chunks, bm25, matrix = self.chunks, self.bm25, self.matrix
        if not chunks:
            return []
        rankings = []
        bm = bm25.scores(tokenize(query))
        rankings.append([i for i in sorted(range(len(bm)), key=lambda i: -bm[i]) if bm[i] > 0][:30])
        if matrix is not None:
            try:
                sims = matrix @ embed([query])[0]
                rankings.append([int(i) for i in sims.argsort()[::-1][:30]])
            except Exception as e:
                log.warning("임베딩 검색 실패, 키워드 검색 결과만 사용: %s", e)
        fused = Counter()
        for ranking in rankings:
            for rank, i in enumerate(ranking):
                fused[i] += 1 / (60 + rank)
        return [{"doc_id": chunks[i]["doc_id"], "title": chunks[i]["title"], "text": chunks[i]["text"],
                 "score": round(s, 4)} for i, s in fused.most_common(k)]


index = Index()


def ingest(title: str, text: str) -> int:
    chunks = chunk_text(text)
    if not chunks:
        raise ValueError("문서에서 텍스트를 찾지 못했습니다. (스캔 PDF는 OCR이 필요합니다)")
    embeddings = None
    if EMBED_MODEL:
        try:
            embeddings = [v.tobytes() for v in embed(chunks)]
        except Exception as e:
            log.warning("임베딩 계산 실패(나중에 다시 시도): %s", e)
    doc_id = store.add_document(title, chunks, embeddings)
    index.rebuild()
    return doc_id
